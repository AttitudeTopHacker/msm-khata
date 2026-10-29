import io
import os
from datetime import datetime
from typing import List, Dict, Any, Optional

import matplotlib
matplotlib.use('Agg')  # Headless backend for server execution
import matplotlib.pyplot as plt

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch, mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, PageBreak, HRFlowable
)
from reportlab.pdfgen import canvas

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas to dynamically compute and print 'Page X of Y' in footer."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        
        # Footer text & line
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.5)
        self.line(36, 30, A4[0] - 36, 30)
        
        self.drawString(36, 20, "MsmKhata Financial Report • Generated via Python Engine • Confidential")
        self.drawRightString(A4[0] - 36, 20, page_str)
        self.restoreState()


class ReportService:
    @staticmethod
    def _create_charts_image(summary: Dict[str, Any], categories: List[Dict[str, Any]]) -> Optional[io.BytesIO]:
        """Generate high-resolution matplotlib charts combining Cashflow, Debt & Category Breakdown."""
        try:
            fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(10, 4.2), dpi=200)
            fig.patch.set_facecolor('#FFFFFF')

            # --- Chart 1: Cashflow & Debt Overview (Grouped Bar Chart) ---
            money_in = summary.get('total_money_in', 0.0)
            money_out = summary.get('total_money_out', 0.0)
            receivable = summary.get('total_receivable', 0.0)
            payable = summary.get('total_payable', 0.0)

            labels = ['Money In\n(Income+Taken)', 'Money Out\n(Expense+Lent)', 'Receivable\n(Lena Hai)', 'Payable / Karz\n(Dena Hai)']
            values = [money_in, money_out, receivable, payable]
            bar_colors = ['#10B981', '#EF4444', '#38BDF8', '#F59E0B']

            bars = ax1.bar(labels, values, color=bar_colors, width=0.55, edgecolor='#CBD5E1', linewidth=0.8)
            ax1.set_title('Cashflow & Market Debt Position (₹)', fontsize=11, fontweight='bold', pad=12, color='#1E293B')
            ax1.grid(axis='y', linestyle='--', alpha=0.4, color='#94A3B8')
            ax1.set_axisbelow(True)
            ax1.tick_params(axis='x', labelsize=8.5, labelcolor='#334155')
            ax1.tick_params(axis='y', labelsize=8.5, labelcolor='#334155')

            # Value labels on top of bars
            for bar in bars:
                height = bar.get_height()
                ax1.annotate(f'₹{int(height):,}',
                             xy=(bar.get_x() + bar.get_width() / 2, height),
                             xytext=(0, 3),  # 3 points vertical offset
                             textcoords="offset points",
                             ha='center', va='bottom', fontsize=8, fontweight='bold', color='#1E293B')

            # --- Chart 2: Category Breakdown (Donut Chart) ---
            if categories and len(categories) > 0:
                cat_names = [c.get('name', 'Other') for c in categories[:5]]
                cat_amounts = [float(c.get('amount', 0.0)) for c in categories[:5]]
                
                # Check if all amounts are zero
                if sum(cat_amounts) > 0:
                    donut_colors = ['#3B82F6', '#10B981', '#F59E0B', '#EC4899', '#8B5CF6']
                    wedges, texts, autotexts = ax2.pie(
                        cat_amounts,
                        labels=cat_names,
                        autopct='%1.0f%%',
                        startangle=140,
                        colors=donut_colors[:len(cat_amounts)],
                        textprops={'fontsize': 8, 'color': '#1E293B'},
                        wedgeprops={'edgecolor': 'white', 'linewidth': 1.5, 'width': 0.45}
                    )
                    for autotext in autotexts:
                        autotext.set_fontsize(7.5)
                        autotext.set_fontweight('bold')
                        autotext.set_color('#FFFFFF')
                    ax2.set_title('Top Expense Categories', fontsize=11, fontweight='bold', pad=12, color='#1E293B')
                else:
                    ax2.text(0.5, 0.5, 'No Expense Categories Yet', ha='center', va='center', color='#94A3B8', fontsize=10)
                    ax2.axis('off')
            else:
                ax2.text(0.5, 0.5, 'No Categorized Expenses', ha='center', va='center', color='#94A3B8', fontsize=10)
                ax2.axis('off')

            plt.tight_layout()
            img_buf = io.BytesIO()
            plt.savefig(img_buf, format='png', bbox_inches='tight', dpi=200)
            plt.close(fig)
            img_buf.seek(0)
            return img_buf
        except Exception as e:
            print(f"Chart generation error: {e}")
            return None

    @classmethod
    def generate_pdf(
        cls,
        account_id: str,
        account_name: str,
        transactions: List[Dict[str, Any]],
        summary: Dict[str, Any],
        persons: Optional[List[Dict[str, Any]]] = None
    ) -> bytes:
        """Build an executive, multi-page financial audit PDF with embedded charts and full passbook ledger."""
        pdf_buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            pdf_buffer,
            pagesize=A4,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=45
        )

        styles = getSampleStyleSheet()
        
        # Custom Typography Styles
        title_style = ParagraphStyle(
            'HeaderTitle',
            parent=styles['Heading1'],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#1E1B4B'),
            fontName='Helvetica-Bold'
        )
        subtitle_style = ParagraphStyle(
            'HeaderSubtitle',
            parent=styles['Normal'],
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor('#64748B'),
            fontName='Helvetica'
        )
        badge_style = ParagraphStyle(
            'BadgeStyle',
            parent=styles['Normal'],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#4F46E5'),
            fontName='Helvetica-Bold',
            alignment=2
        )
        cell_bold = ParagraphStyle(
            'CellBold',
            parent=styles['Normal'],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor('#0F172A'),
            fontName='Helvetica-Bold'
        )
        cell_normal = ParagraphStyle(
            'CellNormal',
            parent=styles['Normal'],
            fontSize=8,
            leading=10.5,
            textColor=colors.HexColor('#334155'),
            fontName='Helvetica'
        )
        cell_green = ParagraphStyle(
            'CellGreen',
            parent=styles['Normal'],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor('#059669'),
            fontName='Helvetica-Bold',
            alignment=2
        )
        cell_red = ParagraphStyle(
            'CellRed',
            parent=styles['Normal'],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor('#DC2626'),
            fontName='Helvetica-Bold',
            alignment=2
        )
        section_heading = ParagraphStyle(
            'SectionHeading',
            parent=styles['Heading2'],
            fontSize=12,
            leading=16,
            textColor=colors.HexColor('#1E293B'),
            fontName='Helvetica-Bold',
            spaceAfter=6
        )

        story = []

        # 1. Header Banner
        now_str = datetime.now().strftime("%d %b %Y, %I:%M %p")
        header_table = Table(
            [
                [
                    Paragraph("<b>MsmKhata Financial Statement</b>", title_style),
                    Paragraph(f"<b>ACCOUNT / BOOK</b><br/>{account_name}", badge_style)
                ],
                [
                    Paragraph("Enterprise Cloud Ledger • ORG: msm.movin 2022", subtitle_style),
                    Paragraph(f"Generated: {now_str}", subtitle_style)
                ]
            ],
            colWidths=[340, 180]
        )
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(header_table)
        story.append(Spacer(1, 10))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#4F46E5'), spaceAfter=14))

        # 2. Key Performance Indicators (Executive Summary Grid)
        net_bal = float(summary.get('net_account_balance', 0.0))
        total_in = float(summary.get('total_money_in', 0.0))
        total_out = float(summary.get('total_money_out', 0.0))
        receivable = float(summary.get('total_receivable', 0.0))
        payable = float(summary.get('total_payable', 0.0))
        net_debt = float(summary.get('net_len_den_position', 0.0))

        kpi_data = [
            [
                Paragraph("<b>NET PASSBOOK BALANCE</b><br/><font size='12' color='%s'><b>₹%s</b></font><br/><font size='7' color='#64748B'>(Total In - Total Out)</font>" %
                          ('#059669' if net_bal >= 0 else '#DC2626', f"{net_bal:,.2f}"), cell_normal),
                Paragraph("<b>TOTAL MONEY IN</b><br/><font size='12' color='#059669'><b>+₹%s</b></font><br/><font size='7' color='#64748B'>Income + Taken</font>" %
                          f"{total_in:,.2f}", cell_normal),
                Paragraph("<b>TOTAL MONEY OUT</b><br/><font size='12' color='#DC2626'><b>-₹%s</b></font><br/><font size='7' color='#64748B'>Expenses + Lent</font>" %
                          f"{total_out:,.2f}", cell_normal),
            ],
            [
                Paragraph("<b>LENA HAI (Receivable)</b><br/><font size='11' color='#0284C7'><b>₹%s</b></font><br/><font size='7' color='#64748B'>Market Udhaar Given</font>" %
                          f"{receivable:,.2f}", cell_normal),
                Paragraph("<b>DENA HAI / KARZ (Payable)</b><br/><font size='11' color='#D97706'><b>₹%s</b></font><br/><font size='7' color='#64748B'>Market Debt Taken</font>" %
                          f"{payable:,.2f}", cell_normal),
                Paragraph("<b>NET MARKET POSITION</b><br/><font size='11' color='%s'><b>₹%s</b></font><br/><font size='7' color='#64748B'>Net Receivable - Payable</font>" %
                          ('#059669' if net_debt >= 0 else '#DC2626', f"{net_debt:,.2f}"), cell_normal),
            ]
        ]

        kpi_table = Table(kpi_data, colWidths=[173, 173, 174])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('LEFTPADDING', (0, 0), (-1, -1), 10),
            ('RIGHTPADDING', (0, 0), (-1, -1), 10),
        ]))
        story.append(kpi_table)
        story.append(Spacer(1, 14))

        # 3. Visual Charts (Cashflow & Donut)
        categories = summary.get('category_breakdown', [])
        chart_buf = cls._create_charts_image(summary, categories)
        if chart_buf:
            story.append(Paragraph("Visual Analytics & Cashflow Trends", section_heading))
            story.append(Image(chart_buf, width=520, height=218))
            story.append(Spacer(1, 14))

        # 4. Detailed Transaction Audit Ledger
        story.append(Paragraph(f"Itemized Transaction Ledger ({len(transactions)} Records)", section_heading))

        headers = [
            Paragraph("<b>#</b>", cell_bold),
            Paragraph("<b>Date</b>", cell_bold),
            Paragraph("<b>Type</b>", cell_bold),
            Paragraph("<b>Party / Category</b>", cell_bold),
            Paragraph("<b>Mode</b>", cell_bold),
            Paragraph("<b>Notes</b>", cell_bold),
            Paragraph("<b>Inflow (+)</b>", cell_bold),
            Paragraph("<b>Outflow (-)</b>", cell_bold),
            Paragraph("<b>Balance</b>", cell_bold)
        ]

        table_rows = [headers]
        
        # Sort chronologically to compute running balance
        sorted_txs = sorted(
            transactions,
            key=lambda x: x.get('transaction_date') or x.get('created_at') or ''
        )

        running_balance = 0.0
        for idx, t in enumerate(sorted_txs, start=1):
            t_type = (t.get('type') or 'OTHER').upper()
            amt = float(t.get('amount', 0.0))
            
            inflow = 0.0
            outflow = 0.0
            
            if t_type in ('INCOME', 'TAKEN'):
                inflow = amt
                running_balance += amt
            elif t_type in ('EXPENSE', 'GIVEN'):
                outflow = amt
                running_balance -= amt

            # Format Date
            t_date_raw = t.get('transaction_date') or t.get('created_at') or ''
            try:
                dt = datetime.fromisoformat(t_date_raw.replace('Z', '+00:00'))
                date_str = dt.strftime("%d %b %y")
            except Exception:
                date_str = t_date_raw[:10] if len(t_date_raw) >= 10 else t_date_raw

            party_or_cat = t.get('person_name') or t.get('category_name') or 'General'
            notes = t.get('notes') or '-'
            if len(notes) > 28:
                notes = notes[:26] + '..'

            # Type Badge Color
            type_label = t_type
            if t_type == 'GIVEN':
                type_label = 'LENT'
            elif t_type == 'TAKEN':
                type_label = 'BORROW'

            inflow_str = f"+₹{inflow:,.0f}" if inflow > 0 else "-"
            outflow_str = f"-₹{outflow:,.0f}" if outflow > 0 else "-"
            bal_str = f"₹{running_balance:,.0f}"

            table_rows.append([
                Paragraph(str(idx), cell_normal),
                Paragraph(date_str, cell_normal),
                Paragraph(type_label, cell_bold),
                Paragraph(party_or_cat, cell_normal),
                Paragraph(t.get('payment_mode') or 'UPI', cell_normal),
                Paragraph(notes, cell_normal),
                Paragraph(inflow_str, cell_green if inflow > 0 else cell_normal),
                Paragraph(outflow_str, cell_red if outflow > 0 else cell_normal),
                Paragraph(bal_str, cell_bold if running_balance >= 0 else cell_red),
            ])

        tx_table = Table(
            table_rows,
            colWidths=[20, 52, 54, 85, 52, 95, 55, 55, 52],
            repeatRows=1
        )
        tx_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (6, 0), (-1, -1), 'RIGHT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))

        story.append(tx_table)

        # Build Document
        doc.build(story, canvasmaker=NumberedCanvas)
        pdf_buffer.seek(0)
        return pdf_buffer.getvalue()

    @classmethod
    def generate_excel(
        cls,
        account_id: str,
        account_name: str,
        transactions: List[Dict[str, Any]],
        summary: Dict[str, Any],
        persons: Optional[List[Dict[str, Any]]] = None
    ) -> bytes:
        """Build an executive, styled multi-tab Excel Workbook (.xlsx) with financial summaries and passbook math."""
        wb = openpyxl.Workbook()

        # Styles definition
        header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        sub_font = Font(name="Calibri", size=10, bold=True, color="1E293B")
        
        green_font = Font(name="Calibri", size=10, bold=True, color="059669")
        red_font = Font(name="Calibri", size=10, bold=True, color="DC2626")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # -------------------------------------------------------------------------
        # SHEET 1: Executive Summary & KPIs
        # -------------------------------------------------------------------------
        ws_sum = wb.active
        ws_sum.title = "Executive Summary"
        ws_sum.views.sheetView[0].showGridLines = True

        ws_sum['A1'] = "MsmKhata Financial Report & Summary"
        ws_sum['A1'].font = Font(name="Calibri", size=16, bold=True, color="1E1B4B")
        ws_sum['A2'] = f"Account / Book: {account_name} (ID: {account_id})"
        ws_sum['A2'].font = Font(name="Calibri", size=11, bold=True, color="4F46E5")
        ws_sum['A3'] = f"Generated On: {datetime.now().strftime('%d %B %Y, %I:%M %p')} • ORG: msm.movin 2022"
        ws_sum['A3'].font = Font(name="Calibri", size=9.5, color="64748B")

        # KPI Table
        kpi_headers = ["Metric Description", "Amount (₹)", "Remarks / Formula"]
        ws_sum.append([])
        ws_sum.append(kpi_headers)
        h_row = ws_sum.max_row
        for col_idx in range(1, 4):
            cell = ws_sum.cell(row=h_row, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center" if col_idx != 1 else "left", vertical="center")

        net_bal = float(summary.get('net_account_balance', 0.0))
        total_in = float(summary.get('total_money_in', 0.0))
        total_out = float(summary.get('total_money_out', 0.0))
        receivable = float(summary.get('total_receivable', 0.0))
        payable = float(summary.get('total_payable', 0.0))
        net_debt = float(summary.get('net_len_den_position', 0.0))

        kpis = [
            ("Real Passbook Available Balance", net_bal, "Real liquid cash/bank funds (Money In - Money Out)"),
            ("Total Money In (Inflow)", total_in, "All Income earned + All Borrowed/Taken received"),
            ("Total Money Out (Outflow)", total_out, "All Expenses paid + All Udhaar/Lent given"),
            ("Market Receivable (Lena Hai)", receivable, "Money owed to you by customers & contacts"),
            ("Market Payable / Karz (Dena Hai)", payable, "Debt/Borrowings you need to repay"),
            ("Net Market Len-Den Position", net_debt, "Receivable minus Payable obligation"),
            ("Total Transactions Logged", len(transactions), "Total count of records in this book")
        ]

        for desc, val, remark in kpis:
            ws_sum.append([desc, val, remark])
            row_idx = ws_sum.max_row
            ws_sum.cell(row=row_idx, column=1).font = bold_font
            val_cell = ws_sum.cell(row=row_idx, column=2)
            val_cell.font = green_font if (isinstance(val, (int, float)) and val > 0 and 'Out' not in desc and 'Payable' not in desc) else (red_font if isinstance(val, (int, float)) and val < 0 else regular_font)
            if isinstance(val, (int, float)) and 'Count' not in desc:
                val_cell.number_format = '₹ #,##0.00'
            ws_sum.cell(row=row_idx, column=3).font = regular_font

            for c in range(1, 4):
                ws_sum.cell(row=row_idx, column=c).border = thin_border

        # -------------------------------------------------------------------------
        # SHEET 2: Full Passbook & Transactions Ledger
        # -------------------------------------------------------------------------
        ws_tx = wb.create_sheet(title="Passbook & Transactions")
        ws_tx.views.sheetView[0].showGridLines = True

        tx_headers = [
            "TX ID", "Date", "Type", "Party / Contact", "Category",
            "Payment Mode", "Notes / Remarks", "Inflow (+) (₹)", "Outflow (-) (₹)", "Running Balance (₹)"
        ]
        ws_tx.append(tx_headers)
        for col_idx in range(1, 11):
            cell = ws_tx.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Sort chronologically for running balance
        sorted_txs = sorted(
            transactions,
            key=lambda x: x.get('transaction_date') or x.get('created_at') or ''
        )

        running_bal = 0.0
        for t in sorted_txs:
            t_type = (t.get('type') or 'OTHER').upper()
            amt = float(t.get('amount', 0.0))
            inflow = 0.0
            outflow = 0.0

            if t_type in ('INCOME', 'TAKEN'):
                inflow = amt
                running_bal += amt
            elif t_type in ('EXPENSE', 'GIVEN'):
                outflow = amt
                running_bal -= amt

            t_date = t.get('transaction_date') or t.get('created_at') or ''
            if len(t_date) >= 10:
                t_date = t_date[:10]

            ws_tx.append([
                t.get('id', '')[:8],
                t_date,
                t_type,
                t.get('person_name') or '-',
                t.get('category_name') or '-',
                t.get('payment_mode') or 'UPI',
                t.get('notes') or '',
                inflow if inflow > 0 else 0.0,
                outflow if outflow > 0 else 0.0,
                running_bal
            ])

            r_idx = ws_tx.max_row
            for c in range(1, 11):
                cell = ws_tx.cell(row=r_idx, column=c)
                cell.font = regular_font
                cell.border = thin_border
            
            # Formatting numeric columns
            in_cell = ws_tx.cell(row=r_idx, column=8)
            in_cell.number_format = '₹ #,##0.00'
            if inflow > 0:
                in_cell.font = green_font

            out_cell = ws_tx.cell(row=r_idx, column=9)
            out_cell.number_format = '₹ #,##0.00'
            if outflow > 0:
                out_cell.font = red_font

            bal_cell = ws_tx.cell(row=r_idx, column=10)
            bal_cell.number_format = '₹ #,##0.00'
            bal_cell.font = green_font if running_bal >= 0 else red_font

        # -------------------------------------------------------------------------
        # SHEET 3: Len-Den Contacts & Parties Ledger
        # -------------------------------------------------------------------------
        if persons and len(persons) > 0:
            ws_parties = wb.create_sheet(title="Parties & Contacts")
            ws_parties.views.sheetView[0].showGridLines = True

            party_headers = ["Party Name", "Contact Phone", "Total Given (Lent)", "Total Taken (Received)", "Net Balance (₹)", "Obligation Status"]
            ws_parties.append(party_headers)
            for col_idx in range(1, 7):
                cell = ws_parties.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for p in persons:
                p_net = float(p.get('net_balance', 0.0))
                status = p.get('status', 'SETTLED')

                ws_parties.append([
                    p.get('name', 'Unknown'),
                    p.get('phone', '-'),
                    float(p.get('total_given', 0.0)),
                    float(p.get('total_taken', 0.0)),
                    p_net,
                    "LENE HAIN (Receivable)" if status == 'TO_RECEIVE' else ("DENE HAIN (Payable)" if status == 'TO_PAY' else "Hisaab Barabar (Settled)")
                ])

                r_idx = ws_parties.max_row
                for c in range(1, 7):
                    cell = ws_parties.cell(row=r_idx, column=c)
                    cell.font = regular_font
                    cell.border = thin_border
                    if c in (3, 4, 5):
                        cell.number_format = '₹ #,##0.00'
                
                ws_parties.cell(row=r_idx, column=5).font = green_font if p_net > 0 else (red_font if p_net < 0 else regular_font)

        # -------------------------------------------------------------------------
        # Auto-fit Column Widths across all sheets
        # -------------------------------------------------------------------------
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        out_buffer = io.BytesIO()
        wb.save(out_buffer)
        out_buffer.seek(0)
        return out_buffer.getvalue()
